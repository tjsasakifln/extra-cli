#!/usr/bin/env python3
"""Canonical weekly operational cycle for Extra Construtora.

Entry point (single):
  make extra-weekly
  python -m scripts.ops.weekly_cycle --strict

Pipeline boundaries:
  collect → process → quality → intelligence → delivery

Exit codes:
  0 — operational cycle valid for consultive use
  1 — technical failure
  2 — completed but not reliable for consultive use
  3 — external block / critical source unavailable
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import time
import traceback
from dataclasses import asdict, dataclass, field
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

_PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

from scripts.collect.run_contract import (  # noqa: E402
    CollectionRun,
    new_collection_id,
    persist_pipeline_run,
)
from scripts.crawl.run_evidence import get_git_meta, new_run_id, sha256_file  # noqa: E402
from scripts.quality.indicator_catalog import catalog_as_list  # noqa: E402

COLLECTOR_VERSION = "weekly-cycle/1.0"
DEFAULT_DSN = os.getenv(
    "LOCAL_DATALAKE_DSN",
    "postgresql://test:test@127.0.0.1:5433/extra_test",
)
# Freshness SLA for reusing previous collection without re-crawl
PNCP_OPP_SLA_HOURS = int(os.getenv("WEEKLY_PNCP_SLA_HOURS", "48"))
CONTRACTS_SLA_HOURS = int(os.getenv("WEEKLY_CONTRACTS_SLA_HOURS", "168"))
DEFAULT_LOOKBACK_DAYS = int(os.getenv("WEEKLY_LOOKBACK_DAYS", "7"))

# Exit codes
EXIT_OK = 0
EXIT_TECH = 1
EXIT_UNRELIABLE = 2
EXIT_BLOCKED = 3


def _utc_now() -> datetime:
    return datetime.now(UTC)


def _iso(dt: datetime | None = None) -> str:
    d = dt or _utc_now()
    return d.replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _resolve_dsn(explicit: str | None) -> str:
    for c in (
        explicit,
        os.getenv("LOCAL_DATALAKE_DSN"),
        os.getenv("DATABASE_URL"),
        DEFAULT_DSN,
    ):
        if c and str(c).strip():
            return str(c).strip()
    raise RuntimeError("No DSN: pass --dsn or set LOCAL_DATALAKE_DSN")


def _connect(dsn: str) -> Any:
    import psycopg2
    import psycopg2.extras

    conn = psycopg2.connect(dsn, cursor_factory=psycopg2.extras.RealDictCursor)
    return conn


def _q(conn: Any, sql: str, params: tuple | list | None = None) -> list[dict[str, Any]]:
    with conn.cursor() as cur:
        cur.execute(sql, params or ())
        rows = cur.fetchall()
        return [dict(r) for r in rows]


def _table_exists(conn: Any, name: str) -> bool:
    rows = _q(
        conn,
        """
        SELECT 1 AS ok FROM information_schema.tables
        WHERE table_schema = 'public' AND table_name = %s
        """,
        (name,),
    )
    return bool(rows)


def _atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + f".{os.getpid()}.tmp")
    tmp.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    os.replace(tmp, path)


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    # stable field order: union of keys preserving first-row order
    fields: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r:
            if k not in seen:
                seen.add(k)
                fields.append(k)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: _csv_cell(r.get(k)) for k in fields})


def _csv_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, default=str)
    return str(v)


@dataclass
class StageResult:
    name: str
    status: str  # ok | warn | fail | blocked | skipped
    detail: dict[str, Any] = field(default_factory=dict)
    error: str | None = None


@dataclass
class WeeklyCycleReport:
    cycle_id: str
    collection_id: str
    started_at: str
    finished_at: str | None = None
    exit_code: int = EXIT_TECH
    git: dict[str, Any] = field(default_factory=dict)
    stages: list[dict[str, Any]] = field(default_factory=list)
    runs: list[dict[str, Any]] = field(default_factory=list)
    freshness: list[dict[str, Any]] = field(default_factory=list)
    source_health: list[dict[str, Any]] = field(default_factory=list)
    gaps: list[dict[str, Any]] = field(default_factory=list)
    products: dict[str, Any] = field(default_factory=dict)
    intelligence: dict[str, Any] = field(default_factory=dict)
    claims_allowed: list[str] = field(default_factory=list)
    claims_forbidden: list[str] = field(default_factory=list)
    limitations: list[str] = field(default_factory=list)
    human_accept: dict[str, Any] = field(default_factory=dict)
    duration_seconds: float = 0.0


# ---------------------------------------------------------------------------
# Stages
# ---------------------------------------------------------------------------


def stage_validate_config(dsn: str) -> StageResult:
    missing = []
    if not dsn:
        missing.append("DSN")
    # openpyxl optional but preferred for Excel
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        missing.append("openpyxl (Excel will degrade to CSV-only note)")
    status = "fail" if "DSN" in missing else ("warn" if missing else "ok")
    return StageResult(
        name="validate_config",
        status=status,
        detail={"dsn_host_set": bool(dsn), "notes": missing},
        error="DSN missing" if status == "fail" else None,
    )


def stage_validate_db(conn: Any) -> StageResult:
    required = [
        "sc_public_entities",
        "opportunity_intel",
        "opportunity_runs",
        "pipeline_runs",
        "pncp_supplier_contracts",
    ]
    present = {t: _table_exists(conn, t) for t in required}
    missing = [t for t, ok in present.items() if not ok]
    if missing:
        return StageResult(
            name="validate_db",
            status="fail",
            detail={"tables": present},
            error=f"missing tables: {missing}",
        )
    uni = _q(
        conn,
        """
        SELECT COUNT(*)::int AS n
        FROM sc_public_entities
        WHERE is_active IS TRUE AND raio_200km IS TRUE
        """,
    )
    n = int((uni[0] or {}).get("n") or 0) if uni else 0
    detail = {"tables": present, "universe_200km": n, "expected_universe": 1093}
    if n != 1093:
        return StageResult(
            name="validate_db",
            status="warn",
            detail=detail,
            error=f"universe_200km={n} != 1093 (scope drift)",
        )
    return StageResult(name="validate_db", status="ok", detail=detail)


def _hours_since(ts: Any) -> float | None:
    if ts is None:
        return None
    if isinstance(ts, str):
        ts = datetime.fromisoformat(ts.replace("Z", "+00:00"))
    if ts.tzinfo is None:
        ts = ts.replace(tzinfo=UTC)
    return max(0.0, (_utc_now() - ts).total_seconds() / 3600.0)


# Only complete, non-error terminal statuses may qualify as fresh.
# ``partial`` is NEVER fresh and NEVER reused_fresh.
_FRESH_OK_RUN_STATUSES = frozenset(
    {
        "completed",
        "success",
        "completed_zero",
        "ok",
    }
)


def classify_opportunity_freshness(
    *,
    status: str | None,
    age_hours: float | None,
    sla_hours: int,
    scope_complete: bool | None = None,
    error_message: str | None = None,
) -> str:
    """Classify source freshness for opportunity collection runs.

    Rules (fail-closed, catalog-aligned):
    - incomplete / partial runs are never ``fresh``
    - only complete success-class statuses within SLA are ``fresh``
    - error or unknown status → ``unreliable``
    """
    st = str(status or "").strip().lower()
    if not st:
        return "never" if age_hours is None else "unreliable"
    if st in {"partial", "running", "interrupted"}:
        return "incomplete"
    if st in {"failed", "error", "blocked", "killed"} or error_message:
        return "unreliable"
    if st not in _FRESH_OK_RUN_STATUSES:
        return "unreliable"
    # Explicit incomplete scope cannot be fresh even if status text is completed
    if scope_complete is False:
        return "incomplete"
    if age_hours is None:
        return "unknown"
    if age_hours <= sla_hours:
        return "fresh"
    return "stale"


def stage_freshness(conn: Any) -> StageResult:
    rows: list[dict[str, Any]] = []
    # PNCP opportunities via opportunity_runs
    opp = _q(
        conn,
        """
        SELECT id, source, status, started_at, finished_at, records_fetched,
               records_new, completion_reason, error_message, scope_complete
        FROM opportunity_runs
        WHERE source LIKE %s
        ORDER BY started_at DESC NULLS LAST
        LIMIT 1
        """,
        ("pncp%",),
    )
    if opp:
        r = opp[0]
        age = _hours_since(r.get("finished_at") or r.get("started_at"))
        st = str(r.get("status") or "")
        scope_complete = r.get("scope_complete")
        level = classify_opportunity_freshness(
            status=st,
            age_hours=age,
            sla_hours=PNCP_OPP_SLA_HOURS,
            scope_complete=scope_complete if scope_complete is not None else None,
            error_message=r.get("error_message"),
        )
        rows.append(
            {
                "source": "pncp_opportunities",
                "level": level,
                "sla_hours": PNCP_OPP_SLA_HOURS,
                "age_hours": round(age, 2) if age is not None else None,
                "last_run_id": r.get("id"),
                "last_status": st,
                "scope_complete": scope_complete,
                "records_fetched": r.get("records_fetched"),
                "indicator": "freshness_source",
            }
        )
    else:
        rows.append(
            {
                "source": "pncp_opportunities",
                "level": "never",
                "sla_hours": PNCP_OPP_SLA_HOURS,
                "age_hours": None,
                "indicator": "freshness_source",
            }
        )

    # Contracts via max(ingested_at) / last_seen
    c = _q(
        conn,
        """
        SELECT MAX(ingested_at) AS last_ingested, COUNT(*)::int AS n
        FROM pncp_supplier_contracts
        WHERE COALESCE(is_active, TRUE)
        """,
    )
    if c and c[0].get("last_ingested"):
        age = _hours_since(c[0]["last_ingested"])
        level = "fresh" if age is not None and age <= CONTRACTS_SLA_HOURS else "stale"
        rows.append(
            {
                "source": "pncp_contracts",
                "level": level,
                "sla_hours": CONTRACTS_SLA_HOURS,
                "age_hours": round(age, 2) if age is not None else None,
                "row_count": c[0].get("n"),
                "indicator": "freshness_source",
                "note": "freshness by max(ingested_at); not a full re-collect this cycle",
            }
        )
    else:
        rows.append(
            {
                "source": "pncp_contracts",
                "level": "never",
                "sla_hours": CONTRACTS_SLA_HOURS,
                "indicator": "freshness_source",
            }
        )

    critical_bad = any(
        r["source"] == "pncp_opportunities"
        and r["level"] in {"never", "unreliable", "incomplete", "unknown"}
        for r in rows
    )
    return StageResult(
        name="freshness",
        status="warn" if critical_bad else "ok",
        detail={"sources": rows},
    )


def _collect_pncp_opportunities(
    conn: Any,
    *,
    collection_id: str,
    dsn: str,
    lookback_days: int,
    force_collect: bool,
    skip_collect: bool,
    freshness_rows: list[dict[str, Any]],
) -> CollectionRun:
    run = CollectionRun.start(
        source="pncp_opportunities",
        collection_id=collection_id,
        collector_version=COLLECTOR_VERSION,
        parameters={
            "lookback_days": lookback_days,
            "force_collect": force_collect,
            "skip_collect": skip_collect,
        },
        period_start=(date.today() - timedelta(days=lookback_days)).isoformat(),
        period_end=date.today().isoformat(),
        mode="incremental",
    )
    pncp_fresh = next(
        (r for r in freshness_rows if r.get("source") == "pncp_opportunities"),
        {},
    )
    # Only a complete, in-SLA run may become reused_fresh.
    # partial / incomplete / unreliable never promote to fresh reuse.
    level = pncp_fresh.get("level")
    truly_fresh = level == "fresh"
    if not force_collect and truly_fresh:
        run.finish(
            records_obtained=int(pncp_fresh.get("records_fetched") or 0),
            records_persisted=int(pncp_fresh.get("records_fetched") or 0),
            request_completed=True,
            scope_complete=True,
            reused_within_sla=True,
            raw_uri=f"db://opportunity_runs/{pncp_fresh.get('last_run_id')}",
            notes=[
                "reused previous COMPLETE PNCP opportunity collection within SLA",
                f"prior_run_status={pncp_fresh.get('last_status')}",
            ],
        )
        try:
            persist_pipeline_run(conn, run)
        except Exception as exc:  # noqa: BLE001
            run.notes.append(f"persist_pipeline_run warn: {exc}")
            conn.rollback()
        return run

    # Explicit skip without fresh complete evidence → partial, never reused_fresh
    if skip_collect:
        run.finish(
            records_obtained=int(pncp_fresh.get("records_fetched") or 0),
            records_persisted=int(pncp_fresh.get("records_fetched") or 0),
            request_completed=True,
            scope_complete=False,
            reused_within_sla=False,
            raw_uri=f"db://opportunity_runs/{pncp_fresh.get('last_run_id')}",
            notes=[
                f"skip_collect with freshness level={level} — not promoted to reused_fresh",
                "partial: lake reused without complete in-SLA collection proof",
            ],
        )
        run.terminal_status = "partial"
        try:
            persist_pipeline_run(conn, run)
        except Exception as exc:  # noqa: BLE001
            run.notes.append(f"persist_pipeline_run warn: {exc}")
            conn.rollback()
        return run

    # Live collect via opportunity_intel crawler
    try:
        from scripts.crawl.pncp_contract import DEFAULT_MODALIDADES
        from scripts.opportunity_intel.crawler_base import CrawlRequest
        from scripts.opportunity_intel.pncp_crawler import PncpOpportunityCrawler

        crawler = PncpOpportunityCrawler(dsn=dsn)
        fetched = 0
        persisted = 0
        rejected = 0
        errors: list[str] = None  # type: ignore[assignment]
        errors = []
        modalidades_ok = 0
        modalidades_total = 0
        try:
            for m in DEFAULT_MODALIDADES:
                modalidades_total += 1
                request = CrawlRequest(
                    source="pncp",
                    date_from=date.today() - timedelta(days=lookback_days),
                    date_to=date.today(),
                    mode="full",
                    limit=None,
                    target=f"modalidade:{m}",
                )
                try:
                    result = crawler.run(request)
                except Exception as exc:  # noqa: BLE001
                    errors.append(f"modalidade:{m}:{exc}")
                    continue
                counts = result.get("counts") or {}
                fetched += int(counts.get("fetched") or 0)
                persisted += int(counts.get("new") or 0) + int(counts.get("updated") or 0)
                st = str(result.get("status") or "")
                if st in {"success", "ok", "completed", "completed_zero", "partial"}:
                    modalidades_ok += 1
                elif result.get("error"):
                    errors.append(f"modalidade:{m}:{result.get('error')}")
        finally:
            crawler.close()

        scope_complete = modalidades_ok == modalidades_total and not errors
        request_completed = modalidades_ok > 0
        err = "; ".join(errors[:5]) if errors else None
        if modalidades_ok == 0 and errors:
            # likely source down
            source_available = not any(
                "timeout" in e.lower() or "connection" in e.lower() or "503" in e or "429" in e
                for e in errors
            )
            # if all failed with HTTP/network → blocked-ish; classify via finish
            run.finish(
                records_obtained=fetched,
                records_persisted=persisted,
                records_rejected=rejected,
                request_completed=False,
                scope_complete=False,
                source_available=source_available,
                error=err or "all modalidades failed",
                notes=[f"modalidades_ok={modalidades_ok}/{modalidades_total}"],
            )
        else:
            run.finish(
                records_obtained=fetched,
                records_persisted=persisted,
                records_rejected=rejected,
                request_completed=request_completed,
                scope_complete=scope_complete,
                source_available=True,
                error=err,
                watermark=f"period:{run.period_start}:{run.period_end}",
                raw_uri="api://pncp.gov.br/api/consulta/v1/contratacoes/proposta",
                notes=[f"modalidades_ok={modalidades_ok}/{modalidades_total}"],
            )
    except Exception as exc:  # noqa: BLE001
        run.finish(
            request_completed=False,
            scope_complete=False,
            source_available=False,
            error=str(exc),
            interrupted=True,
        )

    try:
        persist_pipeline_run(conn, run)
    except Exception as exc:  # noqa: BLE001
        run.notes.append(f"persist_pipeline_run warn: {exc}")
        conn.rollback()
    return run


def _contracts_incremental_run(
    conn: Any,
    *,
    collection_id: str,
    dsn: str,
    days: int = 7,
) -> CollectionRun:
    """Run canonical incremental contracts update (fail-closed on incomplete)."""
    run = CollectionRun.start(
        source="pncp_contracts",
        collection_id=collection_id,
        collector_version=COLLECTOR_VERSION,
        parameters={"strategy": "incremental_via_pilot_runner", "days": days},
        mode="incremental",
    )
    try:
        from scripts.crawl.run_contracts_incremental import main as inc_main

        out = (
            _PROJECT_ROOT
            / "output"
            / "contracts"
            / f"incremental-weekly-{collection_id}.json"
        )
        rc = inc_main(
            [
                "--dsn",
                dsn,
                "--days",
                str(days),
                "--output-json",
                str(out),
                "--checkpoint-dir",
                "data/contracts_checkpoints/weekly_incremental",
                "--reset-checkpoint",
            ]
        )
        payload: dict[str, Any] = {}
        if out.is_file():
            payload = json.loads(out.read_text(encoding="utf-8"))
        totals = payload.get("totals") or {}
        inserted = int(totals.get("inserted") or 0)
        if rc == 0 and str(payload.get("status")) == "success":
            run.finish(
                records_obtained=int(totals.get("fetched") or inserted),
                records_persisted=inserted,
                request_completed=True,
                scope_complete=True,
                reused_within_sla=False,
                raw_uri="api://pncp.gov.br/api/consulta/v1/contratos",
                notes=[
                    "contracts incremental completed this cycle",
                    f"days={days}",
                    f"artifact={out}",
                ],
            )
        else:
            run.finish(
                records_obtained=int(totals.get("fetched") or 0),
                records_persisted=inserted,
                request_completed=False,
                scope_complete=False,
                error=f"incremental_rc={rc} status={payload.get('status')}",
                notes=["contracts incremental incomplete — fail-closed"],
            )
            run.terminal_status = "partial"
    except Exception as exc:  # noqa: BLE001
        run.finish(
            request_completed=False,
            scope_complete=False,
            source_available=False,
            error=str(exc),
            interrupted=True,
        )
    try:
        persist_pipeline_run(conn, run)
    except Exception as exc:  # noqa: BLE001
        run.notes.append(f"persist_pipeline_run warn: {exc}")
        conn.rollback()
    return run


def _contracts_reuse_run(
    conn: Any,
    *,
    collection_id: str,
    freshness_rows: list[dict[str, Any]],
) -> CollectionRun:
    """Contracts: reuse lake data when fresh; stale is partial (fail-closed under --strict)."""
    run = CollectionRun.start(
        source="pncp_contracts",
        collection_id=collection_id,
        collector_version=COLLECTOR_VERSION,
        parameters={"strategy": "reuse_lake_with_freshness_declaration"},
        mode="reuse",
    )
    fr = next((r for r in freshness_rows if r.get("source") == "pncp_contracts"), {})
    n = int(fr.get("row_count") or 0)
    level = fr.get("level")
    if level == "fresh" and n > 0:
        run.finish(
            records_obtained=n,
            records_persisted=n,
            request_completed=True,
            scope_complete=False,  # we did not re-query official API this cycle
            reused_within_sla=True,
            raw_uri="db://pncp_supplier_contracts",
            notes=[
                "contracts not re-crawled; lake rows reused with explicit freshness",
                f"age_hours={fr.get('age_hours')}",
            ],
        )
    elif n > 0:
        run.finish(
            records_obtained=n,
            records_persisted=n,
            request_completed=True,
            scope_complete=False,
            reused_within_sla=False,
            raw_uri="db://pncp_supplier_contracts",
            error=None,
            notes=[
                "contracts lake reused but STALE relative to SLA",
                f"level={level} age_hours={fr.get('age_hours')}",
            ],
        )
        # force partial semantics: not reused_within_sla and scope incomplete
        run.terminal_status = "partial"
        run.notes.append("terminal_status forced to partial due to stale contracts lake")
    else:
        run.finish(
            request_completed=False,
            scope_complete=False,
            source_available=True,
            error="no contracts in lake",
        )
    try:
        persist_pipeline_run(conn, run)
    except Exception as exc:  # noqa: BLE001
        run.notes.append(f"persist_pipeline_run warn: {exc}")
        conn.rollback()
    return run


def stage_quality(conn: Any, runs: list[CollectionRun]) -> StageResult:
    issues: list[str] = []
    # metric separation check — catalog load
    catalog = catalog_as_list()
    if not catalog:
        issues.append("empty indicator catalog")

    # identity safety: sample orgao_cnpj vs fornecedor never promote on name alone
    # (structural check that resolver policy exists)
    try:
        from scripts.entity_identity.pncp_orgao_resolve import pick_match

        # adversarial unit: different root must not match by name
        bad = pick_match(
            "12345678",
            "PREFEITURA MUNICIPAL DEMO",
            [{"cnpj": "99999999000199", "razaoSocial": "PREFEITURA MUNICIPAL DEMO"}],
        )
        if bad is not None:
            issues.append("identity resolver promoted cross-root name match")
    except Exception as exc:  # noqa: BLE001
        issues.append(f"identity check error: {exc}")

    opp_run = next((r for r in runs if r.source == "pncp_opportunities"), None)
    if opp_run and opp_run.terminal_status in {"failure", "blocked"}:
        issues.append(f"pncp_opportunities terminal={opp_run.terminal_status}")

    counts = _q(
        conn,
        """
        SELECT
          COUNT(*) FILTER (WHERE is_active AND status_canonico IN ('open','upcoming'))::int AS open_n,
          COUNT(*) FILTER (WHERE is_active)::int AS active_n,
          COUNT(*) FILTER (WHERE is_active AND ranking = 'GO')::int AS go_n,
          COUNT(*) FILTER (WHERE is_active AND ranking = 'REVIEW')::int AS review_n,
          COUNT(*) FILTER (WHERE is_active AND ranking = 'NO_GO')::int AS nogo_n
        FROM opportunity_intel
        """,
    )
    detail = {
        "catalog_size": len(catalog),
        "opportunity_counts": counts[0] if counts else {},
        "issues": issues,
        "runs": [r.terminal_status for r in runs],
    }
    if any("blocked" == r.terminal_status for r in runs if r.source == "pncp_opportunities"):
        return StageResult(name="quality", status="blocked", detail=detail, error="; ".join(issues) or "blocked")
    if issues:
        return StageResult(name="quality", status="warn", detail=detail, error="; ".join(issues))
    return StageResult(name="quality", status="ok", detail=detail)


# SQL fragment: contracts for Extra weekly pack.
# Require UF=SC (avoids federal CNPJ-8 roots matching nationwide RFB/etc.)
# AND orgao_cnpj_8 ∈ universe raio_200km. Never national ORDER BY alone.
_EXTRA_UNIVERSE_ORGAO = """
    c.uf = 'SC'
    AND c.orgao_cnpj_8 IN (
        SELECT e.cnpj_8
        FROM sc_public_entities e
        WHERE e.is_active IS TRUE
          AND e.raio_200km IS TRUE
          AND e.cnpj_8 IS NOT NULL
          AND LENGTH(TRIM(e.cnpj_8)) = 8
    )
"""


def stage_intelligence(
    conn: Any,
    *,
    limit: int = 50,
    collection_id: str | None = None,
    runs: list[CollectionRun] | None = None,
) -> StageResult:
    runs = runs or []
    run_by_source = {r.source: r for r in runs}
    opp_cycle = run_by_source.get("pncp_opportunities")
    ct_cycle = run_by_source.get("pncp_contracts")

    opps = _q(
        conn,
        """
        SELECT id, source, source_id, numero_controle_pncp, orgao_cnpj, orgao_nome,
               municipio, uf, objeto, modalidade, valor_estimado, valor_semantica,
               status_canonico, ranking, ranking_score, ranking_confianca,
               data_publicacao, data_abertura, data_encerramento,
               link_edital, run_id, crawl_batch_id, proveniencia,
               last_seen_source_run_id, ingested_at, updated_at
        FROM opportunity_intel
        WHERE is_active = TRUE
          AND status_canonico IN ('open', 'upcoming')
          AND (
            uf = 'SC'
            OR LEFT(REGEXP_REPLACE(COALESCE(orgao_cnpj, ''), '[^0-9]', '', 'g'), 8)
               IN (
                 SELECT e.cnpj_8 FROM sc_public_entities e
                 WHERE e.is_active IS TRUE AND e.raio_200km IS TRUE
               )
          )
        ORDER BY
          CASE ranking WHEN 'GO' THEN 0 WHEN 'REVIEW' THEN 1 ELSE 2 END,
          data_encerramento NULLS LAST,
          ranking_score DESC NULLS LAST
        LIMIT %s
        """,
        (limit,),
    )
    # sanitize GO with missing essentials → note only (do not rewrite DB in weekly)
    for o in opps:
        missing = []
        if not o.get("orgao_cnpj"):
            missing.append("orgao_cnpj")
        if not o.get("objeto"):
            missing.append("objeto")
        if o.get("valor_estimado") is None:
            missing.append("valor_estimado")
        o["essential_gaps"] = missing
        if o.get("ranking") == "GO" and missing:
            o["ranking_effective"] = "REVIEW"
            o["ranking_note"] = "GO rebaixado a REVIEW: informações essenciais ausentes"
        else:
            o["ranking_effective"] = o.get("ranking")
            o["ranking_note"] = None
        # value semantics
        o["valor_tipo"] = o.get("valor_semantica") or "estimado_ou_nao_declarado"
        o["valor_nao_e_pago"] = True
        o["cycle_collection_id"] = collection_id
        o["cycle_run_id"] = opp_cycle.run_id if opp_cycle else None
        o["source_record_run_id"] = o.get("run_id") or o.get("last_seen_source_run_id")
        o["scope"] = "extra_sc_or_universe_200km"

    contracts = _q(
        conn,
        # Constant fragment only — not user input (S608 false positive)
        f"""
        SELECT c.contrato_id, c.orgao_cnpj, c.orgao_nome, c.fornecedor_cnpj, c.fornecedor_nome,
               c.objeto_contrato, c.valor_total, c.data_inicio, c.data_fim, c.data_publicacao,
               c.uf, c.municipio, c.source, c.source_id, c.ingested_at, c.source_date_semantics,
               c.orgao_cnpj_8
        FROM pncp_supplier_contracts c
        WHERE COALESCE(c.is_active, TRUE)
          AND {_EXTRA_UNIVERSE_ORGAO}
          AND (
            c.data_publicacao >= CURRENT_DATE - INTERVAL '90 days'
            OR c.data_fim >= CURRENT_DATE - INTERVAL '30 days'
            OR c.data_fim BETWEEN CURRENT_DATE AND CURRENT_DATE + INTERVAL '180 days'
          )
        ORDER BY COALESCE(c.data_publicacao, c.data_fim) DESC NULLS LAST
        LIMIT %s
        """,  # noqa: S608
        (limit,),
    )
    for c in contracts:
        c["valor_tipo"] = "valor_contratado"
        c["valor_nao_e_pago"] = True
        c["valor_note"] = "valor_total é contratado/homologado na fonte — não pagar/medido"
        c["cycle_collection_id"] = collection_id
        c["cycle_run_id"] = ct_cycle.run_id if ct_cycle else None
        # source_id is a record identifier, not an execution run — do not mislabel it
        c["source_record_run_id"] = None
        c["source_record_id"] = c.get("source_id") or c.get("contrato_id")
        c["scope"] = "extra_universe_200km"
        c["normalized_table"] = "pncp_supplier_contracts"

    competitors = _q(
        conn,
        f"""
        SELECT c.fornecedor_cnpj, c.fornecedor_nome,
               COUNT(*)::int AS n_contratos,
               SUM(c.valor_total)::numeric AS soma_valor_contratado,
               COUNT(DISTINCT c.orgao_cnpj)::int AS n_orgaos
        FROM pncp_supplier_contracts c
        WHERE COALESCE(c.is_active, TRUE)
          AND {_EXTRA_UNIVERSE_ORGAO}
          AND c.fornecedor_cnpj IS NOT NULL
          AND COALESCE(c.data_publicacao, c.data_inicio, c.ingested_at)
              >= CURRENT_DATE - INTERVAL '365 days'
        GROUP BY c.fornecedor_cnpj, c.fornecedor_nome
        ORDER BY n_contratos DESC, soma_valor_contratado DESC NULLS LAST
        LIMIT %s
        """,  # noqa: S608
        (limit,),
    )
    for c in competitors:
        c["valor_tipo"] = "soma_valor_contratado_nao_pago"
        c["valor_nao_e_pago"] = True
        c["cycle_collection_id"] = collection_id
        c["cycle_run_id"] = ct_cycle.run_id if ct_cycle else None
        c["scope"] = "extra_universe_200km"
        c["normalized_table"] = "pncp_supplier_contracts_agg"

    orgaos = _q(
        conn,
        """
        SELECT orgao_cnpj, orgao_nome, uf, municipio, COUNT(*)::int AS n_opp
        FROM opportunity_intel
        WHERE is_active AND status_canonico IN ('open','upcoming')
          AND (
            uf = 'SC'
            OR LEFT(REGEXP_REPLACE(COALESCE(orgao_cnpj, ''), '[^0-9]', '', 'g'), 8)
               IN (
                 SELECT e.cnpj_8 FROM sc_public_entities e
                 WHERE e.is_active IS TRUE AND e.raio_200km IS TRUE
               )
          )
        GROUP BY orgao_cnpj, orgao_nome, uf, municipio
        ORDER BY n_opp DESC
        LIMIT %s
        """,
        (limit,),
    )
    for o in orgaos:
        o["cycle_collection_id"] = collection_id
        o["scope"] = "extra_sc_or_universe_200km"

    return StageResult(
        name="intelligence",
        status="ok" if opps or contracts else "warn",
        detail={
            "opportunities": opps,
            "contracts": contracts,
            "competitors": competitors,
            "orgaos": orgaos,
            "counts": {
                "opportunities": len(opps),
                "contracts": len(contracts),
                "competitors": len(competitors),
                "orgaos": len(orgaos),
            },
            "scope": "extra_universe_200km",
        },
    )


def _build_claims_catalog(
    intel: dict[str, Any],
    runs: list[CollectionRun],
    freshness: list[dict[str, Any]],
    *,
    collection_id: str | None = None,
) -> list[dict[str, Any]]:
    """Provenance catalog: source → collection → cycle run → normalized → rule → product.

    Material rows (opportunities, contracts, competitors) must carry this cycle's
    ``collection_id`` and ``cycle_run_id``. Source-record run ids are preserved
    separately and must not replace the cycle collection link.
    """
    claims: list[dict[str, Any]] = []
    run_by_source = {r.source: r for r in runs}
    opp_cycle = run_by_source.get("pncp_opportunities")
    ct_cycle = run_by_source.get("pncp_contracts")
    cid = collection_id or (runs[0].collection_id if runs else None)

    for o in intel.get("opportunities") or []:
        claims.append(
            {
                "claim_id": f"opp-{o.get('id')}",
                "kind": "opportunity",
                "statement": (
                    f"Oportunidade {o.get('numero_controle_pncp') or o.get('source_id')} "
                    f"orgão={o.get('orgao_nome')} ranking={o.get('ranking_effective')} "
                    f"valor_estimado={o.get('valor_estimado')} ({o.get('valor_tipo')})"
                ),
                "source": o.get("source") or "pncp",
                "collection_id": o.get("cycle_collection_id") or cid,
                "cycle_run_id": o.get("cycle_run_id")
                or (opp_cycle.run_id if opp_cycle else None),
                "source_record_run_id": o.get("source_record_run_id")
                or o.get("run_id")
                or o.get("last_seen_source_run_id"),
                "normalized_table": "opportunity_intel",
                "normalized_id": o.get("id"),
                "rule": "opportunity_intel.ranking + status_canonico + extra_scope",
                "product": "opportunities.csv / executive_summary.md",
            }
        )

    for c in intel.get("contracts") or []:
        claims.append(
            {
                "claim_id": f"contract-{c.get('contrato_id') or c.get('source_id')}",
                "kind": "contract",
                "statement": (
                    f"Contrato órgão={c.get('orgao_nome')} fornecedor={c.get('fornecedor_nome')} "
                    f"valor_contratado={c.get('valor_total')} ({c.get('valor_tipo')})"
                ),
                "source": c.get("source") or "pncp_contracts",
                "collection_id": c.get("cycle_collection_id") or cid,
                "cycle_run_id": c.get("cycle_run_id")
                or (ct_cycle.run_id if ct_cycle else None),
                "source_record_run_id": c.get("source_record_run_id"),  # often null for lake reuse
                "source_record_id": c.get("source_record_id")
                or c.get("source_id")
                or c.get("contrato_id"),
                "normalized_table": "pncp_supplier_contracts",
                "normalized_id": c.get("contrato_id") or c.get("source_id"),
                "rule": "extra_universe_200km + valor_contratado_not_paid",
                "product": "contracts.csv / executive_summary.md",
                "scope": c.get("scope") or "extra_universe_200km",
            }
        )

    for c in intel.get("competitors") or []:
        claims.append(
            {
                "claim_id": f"competitor-{c.get('fornecedor_cnpj')}",
                "kind": "competitor",
                "statement": (
                    f"Concorrente {c.get('fornecedor_nome')} cnpj={c.get('fornecedor_cnpj')} "
                    f"n_contratos={c.get('n_contratos')} "
                    f"soma_valor_contratado={c.get('soma_valor_contratado')} "
                    f"({c.get('valor_tipo')})"
                ),
                "source": "pncp_contracts",
                "collection_id": c.get("cycle_collection_id") or cid,
                "cycle_run_id": c.get("cycle_run_id")
                or (ct_cycle.run_id if ct_cycle else None),
                "source_record_run_id": None,
                "normalized_table": "pncp_supplier_contracts_agg",
                "normalized_id": c.get("fornecedor_cnpj"),
                "rule": "extra_universe_200km competitor aggregation 365d",
                "product": "competitors.csv / executive_summary.md",
                "scope": c.get("scope") or "extra_universe_200km",
            }
        )

    for r in runs:
        claims.append(
            {
                "claim_id": f"run-{r.run_id}",
                "kind": "collection_run",
                "statement": (
                    f"Coleta {r.source} terminal={r.terminal_status} "
                    f"obtained={r.records_obtained} persisted={r.records_persisted}"
                ),
                "source": r.source,
                "collection_id": r.collection_id,
                "cycle_run_id": r.run_id,
                "source_record_run_id": None,
                "normalized_table": "pipeline_runs",
                "normalized_id": r.run_id,
                "rule": "scripts.collect.run_contract",
                "product": "manifest.json",
            }
        )
    for f in freshness:
        claims.append(
            {
                "claim_id": f"fresh-{f.get('source')}",
                "kind": "freshness",
                "statement": (
                    f"Freshness {f.get('source')}={f.get('level')} "
                    f"age_hours={f.get('age_hours')} sla={f.get('sla_hours')}"
                ),
                "source": f.get("source"),
                "collection_id": cid,
                "cycle_run_id": None,
                "indicator": f.get("indicator"),
                "normalized_table": None,
                "normalized_id": None,
                "rule": "freshness_source",
                "product": "source_health.csv / executive_summary.md",
            }
        )
    return claims


def stage_delivery(
    out_dir: Path,
    report: WeeklyCycleReport,
    intel: dict[str, Any],
    runs: list[CollectionRun],
    freshness: list[dict[str, Any]],
    gaps: list[dict[str, Any]],
) -> StageResult:
    out_dir.mkdir(parents=True, exist_ok=True)
    products: dict[str, Any] = {}

    opps = intel.get("opportunities") or []
    contracts = intel.get("contracts") or []
    competitors = intel.get("competitors") or []
    orgaos = intel.get("orgaos") or []

    # CSVs
    paths = {
        "opportunities_csv": out_dir / "opportunities.csv",
        "contracts_csv": out_dir / "contracts.csv",
        "competitors_csv": out_dir / "competitors.csv",
        "orgaos_csv": out_dir / "orgaos.csv",
        "source_health_csv": out_dir / "source_health.csv",
        "gaps_csv": out_dir / "gaps.csv",
        "claims_csv": out_dir / "claims_provenance.csv",
    }
    _write_csv(paths["opportunities_csv"], opps)
    _write_csv(paths["contracts_csv"], contracts)
    _write_csv(paths["competitors_csv"], competitors)
    _write_csv(paths["orgaos_csv"], orgaos)
    _write_csv(paths["source_health_csv"], freshness)
    _write_csv(paths["gaps_csv"], gaps)
    claims = _build_claims_catalog(
        intel,
        runs,
        freshness,
        collection_id=report.collection_id,
    )
    _write_csv(paths["claims_csv"], claims)

    # Excel
    xlsx_path = out_dir / "extra_weekly_pack.xlsx"
    excel_ok = False
    excel_note = ""
    try:
        from openpyxl import Workbook

        wb = Workbook()
        meta = wb.active
        meta.title = "Metadados"
        meta.append(["key", "value"])
        for k, v in [
            ("cycle_id", report.cycle_id),
            ("collection_id", report.collection_id),
            ("started_at", report.started_at),
            ("git_sha", (report.git or {}).get("git_sha")),
            ("valor_semantica", "estimado≠homologado≠pago"),
            ("pdf_status", "RESIDUAL — não gerado como produto operacional"),
        ]:
            meta.append([k, str(v)])
        for name, rows in [
            ("Oportunidades", opps),
            ("Contratos", contracts),
            ("Concorrentes", competitors),
            ("Orgaos", orgaos),
            ("SourceHealth", freshness),
            ("Gaps", gaps),
            ("Limitacoes", [{"limitacao": x} for x in report.limitations]),
        ]:
            ws = wb.create_sheet(name)
            if not rows:
                ws.append(["(vazio)"])
                continue
            headers = list(rows[0].keys())
            ws.append(headers)
            for r in rows:
                ws.append([_csv_cell(r.get(h)) for h in headers])
        wb.save(xlsx_path)
        excel_ok = True
    except Exception as exc:  # noqa: BLE001
        excel_note = str(exc)
        xlsx_path = out_dir / "extra_weekly_pack.excel_failed.txt"
        xlsx_path.write_text(f"Excel generation failed: {exc}\n", encoding="utf-8")

    # Executive markdown
    md_path = out_dir / "executive_summary.md"
    go_n = sum(1 for o in opps if o.get("ranking_effective") == "GO")
    review_n = sum(1 for o in opps if o.get("ranking_effective") == "REVIEW")
    nogo_n = sum(1 for o in opps if o.get("ranking_effective") == "NO_GO")
    lines = [
        f"# Pacote semanal Extra Construtora — {report.cycle_id}",
        "",
        f"- **Gerado em:** {report.started_at}",
        f"- **Collection ID:** `{report.collection_id}`",
        f"- **Git:** `{(report.git or {}).get('git_sha') or 'unknown'}`",
        "- **Exit code previsto:** ver manifest",
        "",
        "## Resumo executivo",
        "",
        f"Este pacote lista **{len(opps)}** oportunidades abertas/upcoming, "
        f"**{len(contracts)}** contratos recentes/relevantes, "
        f"**{len(competitors)}** concorrentes observáveis (top) e "
        f"**{len(orgaos)}** órgãos associados nas oportunidades.",
        "",
        f"Ranking efetivo: **GO={go_n}**, **REVIEW={review_n}**, **NO_GO={nogo_n}**.",
        "",
        "> Scores **não** são probabilidades de vitória.",
        ">",
        "> Valores de contrato são **contratados**, não pagos.",
        "",
        "## Freshness / saúde das fontes",
        "",
    ]
    for f in freshness:
        lines.append(
            f"- `{f.get('source')}`: **{f.get('level')}** "
            f"(age_h={f.get('age_hours')}, SLA={f.get('sla_hours')}h)"
            + (f" — {f.get('note')}" if f.get("note") else "")
        )
    lines += ["", "## Coletas deste ciclo", ""]
    for r in runs:
        lines.append(
            f"- `{r.source}` run `{r.run_id}` → **{r.terminal_status}** "
            f"(obtidos={r.records_obtained}, persistidos={r.records_persisted})"
        )
        if r.terminal_error:
            lines.append(f"  - erro: {r.terminal_error}")
        for n in r.notes[:3]:
            lines.append(f"  - nota: {n}")

    lines += ["", "## Top oportunidades (até 15)", ""]
    lines.append("| id | ranking | órgão | objeto | valor_estimado | prazo | fonte |")
    lines.append("|---:|---|---|---|---:|---|---|")
    for o in opps[:15]:
        obj = (o.get("objeto") or "")[:80].replace("|", "/")
        org = (o.get("orgao_nome") or "")[:40].replace("|", "/")
        lines.append(
            f"| {o.get('id')} | {o.get('ranking_effective')} | {org} | {obj} | "
            f"{o.get('valor_estimado')} | {o.get('data_encerramento') or '—'} | {o.get('source')} |"
        )

    lines += ["", "## Contratos (amostra)", ""]
    lines.append("| órgão | fornecedor | valor_contratado | fim |")
    lines.append("|---|---|---:|---|")
    for c in contracts[:10]:
        lines.append(
            f"| {(c.get('orgao_nome') or '')[:40]} | {(c.get('fornecedor_nome') or '')[:40]} | "
            f"{c.get('valor_total')} | {c.get('data_fim') or '—'} |"
        )

    lines += ["", "## Gaps conhecidos", ""]
    if gaps:
        for g in gaps:
            lines.append(f"- {g.get('gap')}: {g.get('detail')}")
    else:
        lines.append("- (nenhum gap estrutural adicional registrado neste ciclo)")

    lines += ["", "## Limitações", ""]
    for lim in report.limitations:
        lines.append(f"- {lim}")
    lines += [
        "",
        "## Aceite humano",
        "",
        "Status: **PENDING_HUMAN** (Tiago). "
        "Ausência de manifestação **não** é aceite.",
        "",
        "Revisar no mínimo: resumo, oportunidades, amostra de contratos, "
        "concorrentes, valores e limitações.",
        "",
        "## PDF",
        "",
        "**RESIDUAL:** PDF operacional multi-página real não é gate deste ciclo. "
        "Produto canônico: Markdown + Excel + CSV.",
        "",
    ]
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    # Product checksums EXCLUDE the manifest (no self-referential hash).
    # Written to checksums.json; manifest only references that external file.
    checksums: dict[str, Any] = {}
    product_paths: dict[str, Path] = {
        **{k: Path(v) for k, v in paths.items()},
        "executive_md": md_path,
        "excel": xlsx_path,
    }
    for label, pth in product_paths.items():
        if pth.exists() and pth.is_file():
            checksums[label] = {
                "path": str(
                    pth.relative_to(_PROJECT_ROOT)
                    if pth.is_relative_to(_PROJECT_ROOT)
                    else pth
                ),
                "sha256": sha256_file(pth),
                "bytes": pth.stat().st_size,
            }

    checksums_path = out_dir / "checksums.json"
    _atomic_json(
        checksums_path,
        {
            "schema": "extra-weekly-checksums/1.0",
            "cycle_id": report.cycle_id,
            "collection_id": report.collection_id,
            "note": "Hashes of product artifacts only — does not include manifest.json",
            "artifacts": checksums,
        },
    )
    checksums_file_meta = {
        "path": str(
            checksums_path.relative_to(_PROJECT_ROOT)
            if checksums_path.is_relative_to(_PROJECT_ROOT)
            else checksums_path
        ),
        "sha256": sha256_file(checksums_path),
        "bytes": checksums_path.stat().st_size if checksums_path.exists() else 0,
    }

    # Excel is part of the product contract — required for delivery=ok
    if excel_ok and md_path.exists() and paths["claims_csv"].exists():
        delivery_status = "ok"
    elif not excel_ok:
        delivery_status = "fail"
        excel_note = excel_note or "Excel generation failed — required for delivery=ok"
    else:
        delivery_status = "fail"

    products = {
        "executive_md": str(md_path),
        "excel": str(xlsx_path),
        "excel_ok": excel_ok,
        "excel_note": excel_note,
        "pdf": None,
        "pdf_status": "RESIDUAL_NOT_GENERATED",
        "csvs": {k: str(v) for k, v in paths.items()},
        "checksums_file": str(checksums_path),
        "checksums_file_meta": checksums_file_meta,
        "product_checksums": checksums,
        "claims_count": len(claims),
    }
    return StageResult(
        name="delivery",
        status=delivery_status,
        detail=products,
        error=None if delivery_status == "ok" else (excel_note or "delivery incomplete"),
    )


def _default_limitations(runs: list[CollectionRun], freshness: list[dict[str, Any]]) -> list[str]:
    lim = [
        "Este pacote não declara LOCAL_READY, cobertura operacional 95% nem recall independente.",
        "Ranking GO/REVIEW/NO_GO é triagem interna, não probabilidade calibrada.",
        "valor_estimado ≠ valor_homologado ≠ valor pago/medido.",
        "PDF multi-página real permanece residual nesta campanha.",
        "Contratos no ciclo semanal reutilizam o lake com declaração de freshness "
        "(re-coleta completa de 499k+ linhas está fora do orçamento do ciclo).",
        "Universo canônico = entidades raio 200 km (meta 1093).",
    ]
    for r in runs:
        if r.terminal_status == "partial":
            lim.append(f"Coleta parcial em {r.source}: {r.terminal_error or r.notes}")
        if r.terminal_status == "reused_fresh":
            lim.append(f"Fonte {r.source} reutilizada dentro do SLA (sem nova chamada oficial).")
        if r.terminal_status in {"failure", "blocked"}:
            lim.append(f"Fonte {r.source} em estado {r.terminal_status}.")
    for f in freshness:
        if f.get("level") in {"stale", "never", "unreliable"}:
            lim.append(f"Freshness {f.get('source')}={f.get('level')}.")
    return lim


def _default_gaps(conn: Any, intel: dict[str, Any]) -> list[dict[str, Any]]:
    gaps: list[dict[str, Any]] = []
    go_n = sum(1 for o in (intel.get("opportunities") or []) if o.get("ranking_effective") == "GO")
    if go_n == 0:
        gaps.append(
            {
                "gap": "no_GO_rankings",
                "detail": "Nenhuma oportunidade com ranking efetivo GO — revisar perfil Extra / fatores",
            }
        )
    missing_cnpj = sum(1 for o in (intel.get("opportunities") or []) if not o.get("orgao_cnpj"))
    if missing_cnpj:
        gaps.append(
            {
                "gap": "opportunities_missing_orgao_cnpj",
                "detail": f"{missing_cnpj} oportunidades sem orgao_cnpj",
            }
        )
    # editais coverage not claimed
    gaps.append(
        {
            "gap": "editais_coverage_below_95",
            "detail": "Cobertura de editais permanece abaixo de 95% — não claim nesta campanha",
        }
    )
    gaps.append(
        {
            "gap": "recall_independent_unproven",
            "detail": "Recall independente estratificado não comprovado",
        }
    )
    if _table_exists(conn, "official_acts"):
        n = _q(conn, "SELECT COUNT(*)::int AS n FROM official_acts")
        if n and int(n[0].get("n") or 0) == 0:
            gaps.append(
                {
                    "gap": "official_acts_empty",
                    "detail": "Tabela official_acts vazia no lake local",
                }
            )
    return gaps


def compute_exit_code(
    stages: list[StageResult],
    runs: list[CollectionRun],
    *,
    strict: bool = True,
) -> int:
    """Exit code policy.

    EXIT_OK only when critical collection is complete (success / success_zero /
    reused_fresh of a complete prior run) AND delivery is ok (incl. Excel).

    ``partial`` never yields EXIT_OK — even outside strict mode.
    ``strict`` additionally fails closed on delivery/artifact defects that
    non-strict might still surface as products with EXIT_UNRELIABLE.
    """
    if any(s.status == "fail" for s in stages if s.name in {"validate_config", "validate_db"}):
        return EXIT_TECH
    opp = next((r for r in runs if r.source == "pncp_opportunities"), None)
    if opp and opp.terminal_status == "blocked":
        return EXIT_BLOCKED
    if any(s.status == "blocked" for s in stages):
        return EXIT_BLOCKED
    if opp and opp.terminal_status == "failure":
        return EXIT_UNRELIABLE
    # Critical partial is never consultively OK
    if opp and opp.terminal_status == "partial":
        return EXIT_UNRELIABLE

    delivery = next((s for s in stages if s.name == "delivery"), None)
    quality = next((s for s in stages if s.name == "quality"), None)
    intel = next((s for s in stages if s.name == "intelligence"), None)

    if delivery and delivery.status == "fail":
        # Missing Excel / incomplete pack is technical delivery failure under strict;
        # still non-zero outside strict (unreliable for consultive use).
        return EXIT_TECH if strict else EXIT_UNRELIABLE

    if quality and quality.status == "blocked":
        return EXIT_BLOCKED

    if strict:
        if delivery is None or delivery.status != "ok":
            return EXIT_UNRELIABLE
        d = delivery.detail or {}
        if not d.get("excel_ok"):
            return EXIT_UNRELIABLE
        if not d.get("checksums_file") and not d.get("product_checksums"):
            return EXIT_UNRELIABLE
        # Contracts section is part of the full consultative package: stale /
        # missing / partial contracts must not yield EXIT_OK under --strict.
        ct = next((r for r in runs if r.source == "pncp_contracts"), None)
        if ct is None or ct.terminal_status not in {
            "success",
            "success_zero",
            "reused_fresh",
        }:
            return EXIT_UNRELIABLE

    if (
        intel
        and intel.status == "ok"
        and delivery
        and delivery.status == "ok"
        and opp
        and opp.terminal_status in {"success", "success_zero", "reused_fresh"}
    ):
        counts = (intel.detail or {}).get("counts") or {}
        if counts.get("opportunities", 0) == 0 and opp.terminal_status not in {
            "success_zero",
            "reused_fresh",
        }:
            return EXIT_UNRELIABLE
        return EXIT_OK
    return EXIT_UNRELIABLE


def run_weekly_cycle(
    *,
    dsn: str | None = None,
    output_dir: Path | None = None,
    strict: bool = True,
    force_collect: bool = False,
    skip_collect: bool = False,
    lookback_days: int = DEFAULT_LOOKBACK_DAYS,
    limit: int = 50,
    offline: bool = False,
    contracts_incremental: bool = False,
    contracts_incremental_days: int = 7,
) -> WeeklyCycleReport:
    """Execute the weekly cycle. `offline` skips live collect (tests).

    When ``contracts_incremental`` is true, run the canonical contracts
    incremental update before packaging (preferred for full consultative packs).
    """
    t0 = time.monotonic()
    cycle_id = new_run_id("weekly")
    collection_id = new_collection_id("extra-weekly")
    git = get_git_meta()
    report = WeeklyCycleReport(
        cycle_id=cycle_id,
        collection_id=collection_id,
        started_at=_iso(),
        git=git,
        human_accept={
            "status": "PENDING_HUMAN",
            "owner": "Tiago",
            "required_reviews": [
                "executive_summary",
                "opportunities",
                "contracts_sample",
                "competitors",
                "values_semantics",
                "limitations",
            ],
            "note": "Ausência de manifestação não é aceite",
        },
        claims_allowed=[
            "Pacote semanal gerado por make extra-weekly / scripts.ops.weekly_cycle",
            "Oportunidades e contratos com proveniência de run/collection quando disponível",
            "Freshness declarado por fonte com SLA",
            "Valores com semântica explícita (estimado vs contratado; nunca como pago)",
        ],
        claims_forbidden=[
            "LOCAL_READY",
            "PRE_VPS_FINAL_READY",
            "VPS_OPERATIONAL",
            "PROJECT_DONE",
            "cobertura operacional 95%",
            "recall independente 95%",
            "proxy de contratos = cobertura completa",
            "score = probabilidade",
            "valor contratado = valor pago",
            "PDF fixture como produto operacional",
        ],
    )

    resolved = _resolve_dsn(dsn)
    stages: list[StageResult] = []
    runs: list[CollectionRun] = []

    sc = stage_validate_config(resolved)
    stages.append(sc)
    if sc.status == "fail":
        report.stages = [asdict(s) for s in stages]
        report.exit_code = EXIT_TECH
        report.finished_at = _iso()
        report.duration_seconds = round(time.monotonic() - t0, 2)
        return report

    try:
        conn = _connect(resolved)
    except Exception as exc:  # noqa: BLE001
        stages.append(
            StageResult(name="validate_db", status="fail", error=str(exc))
        )
        report.stages = [asdict(s) for s in stages]
        report.exit_code = EXIT_TECH
        report.finished_at = _iso()
        report.duration_seconds = round(time.monotonic() - t0, 2)
        report.limitations = [f"DB connect failed: {exc}"]
        return report

    try:
        sdb = stage_validate_db(conn)
        stages.append(sdb)
        if sdb.status == "fail":
            report.exit_code = EXIT_TECH
            report.stages = [asdict(s) for s in stages]
            report.finished_at = _iso()
            report.duration_seconds = round(time.monotonic() - t0, 2)
            return report

        sf = stage_freshness(conn)
        stages.append(sf)
        freshness_rows = list((sf.detail or {}).get("sources") or [])
        report.freshness = freshness_rows

        # collect
        if offline:
            r_opp = CollectionRun.start(
                source="pncp_opportunities",
                collection_id=collection_id,
                collector_version=COLLECTOR_VERSION,
                mode="offline_test",
            )
            r_opp.finish(
                records_obtained=0,
                records_persisted=0,
                request_completed=True,
                scope_complete=True,
                reused_within_sla=True,
                notes=["offline test mode — no network"],
            )
        else:
            r_opp = _collect_pncp_opportunities(
                conn,
                collection_id=collection_id,
                dsn=resolved,
                lookback_days=lookback_days,
                force_collect=force_collect,
                skip_collect=skip_collect,
                freshness_rows=freshness_rows,
            )
        runs.append(r_opp)
        if contracts_incremental and not offline:
            r_ct = _contracts_incremental_run(
                conn,
                collection_id=collection_id,
                dsn=resolved,
                days=contracts_incremental_days,
            )
        else:
            r_ct = _contracts_reuse_run(
                conn, collection_id=collection_id, freshness_rows=freshness_rows
            )
        runs.append(r_ct)
        collect_status = "ok"
        if r_opp.terminal_status == "blocked":
            collect_status = "blocked"
        elif r_opp.terminal_status in {"failure"}:
            collect_status = "fail"
        elif r_opp.terminal_status == "partial":
            collect_status = "warn"
        elif r_opp.terminal_status not in {"success", "success_zero", "reused_fresh"}:
            collect_status = "fail"
        stages.append(
            StageResult(
                name="collect",
                status=collect_status,
                detail={"runs": [r.to_dict() for r in runs]},
            )
        )
        stages.append(
            StageResult(
                name="process",
                status="ok",
                detail={
                    "note": "PNCP opportunity crawler normalizes into opportunity_intel; "
                    "contracts already canonical in pncp_supplier_contracts"
                },
            )
        )

        sq = stage_quality(conn, runs)
        stages.append(sq)
        si = stage_intelligence(
            conn,
            limit=limit,
            collection_id=collection_id,
            runs=runs,
        )
        stages.append(si)
        intel = si.detail or {}
        report.intelligence = {
            "counts": intel.get("counts"),
            # do not dump full rows into report root (size); products hold them
        }
        gaps = _default_gaps(conn, intel)
        report.gaps = gaps
        report.limitations = _default_limitations(runs, freshness_rows)
        report.source_health = freshness_rows

        out = output_dir or (
            _PROJECT_ROOT / "output" / "weekly" / cycle_id
        )
        report.limitations = _default_limitations(runs, freshness_rows)
        sd = stage_delivery(out, report, intel, runs, freshness_rows, gaps)
        stages.append(sd)
        report.products = sd.detail or {}
        report.runs = [r.to_dict() for r in runs]
        report.stages = [asdict(s) for s in stages]
        report.exit_code = compute_exit_code(stages, runs, strict=strict)

        # Write manifest ONCE — no self-hash. Integrity of products is in checksums.json.
        report.finished_at = _iso()
        report.duration_seconds = round(time.monotonic() - t0, 2)
        manifest_path = out / "manifest.json"
        report.products["manifest"] = str(manifest_path)
        report.products["manifest_integrity"] = (
            "product hashes live in checksums.json; manifest is not self-hashed"
        )
        _atomic_json(manifest_path, asdict(report))

        return report
    finally:
        try:
            conn.close()
        except Exception as close_exc:  # noqa: BLE001
            sys.stderr.write(f"weekly_cycle: conn.close warn: {close_exc}\n")


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="weekly_cycle",
        description="Ciclo operacional semanal canônico — Extra Construtora",
    )
    p.add_argument("--dsn", default=None, help="PostgreSQL DSN")
    p.add_argument(
        "--output-dir",
        default=None,
        help="Diretório de saída (default: output/weekly/<cycle_id>)",
    )
    p.add_argument(
        "--strict",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Modo estrito consultivo (default: true)",
    )
    p.add_argument(
        "--force-collect",
        action="store_true",
        help="Força recoleta PNCP mesmo se fresca",
    )
    p.add_argument(
        "--skip-collect",
        action="store_true",
        help="Não chama API; reutiliza lake (ainda declara freshness)",
    )
    p.add_argument("--lookback-days", type=int, default=DEFAULT_LOOKBACK_DAYS)
    p.add_argument("--limit", type=int, default=50, help="Limite por lista de inteligência")
    p.add_argument(
        "--offline",
        action="store_true",
        help="Modo teste sem rede (não usar para pacote real)",
    )
    p.add_argument(
        "--contracts-incremental",
        action=argparse.BooleanOptionalAction,
        default=False,
        help="Executa atualização incremental de contratos no ciclo (recomendado para pack consultivo integral)",
    )
    p.add_argument(
        "--contracts-incremental-days",
        type=int,
        default=7,
        help="Janela lookback do incremental de contratos (default 7)",
    )
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    out = Path(args.output_dir) if args.output_dir else None
    try:
        report = run_weekly_cycle(
            dsn=args.dsn,
            output_dir=out,
            strict=bool(args.strict),
            force_collect=bool(args.force_collect),
            skip_collect=bool(args.skip_collect),
            lookback_days=int(args.lookback_days),
            limit=int(args.limit),
            offline=bool(args.offline),
            contracts_incremental=bool(args.contracts_incremental),
            contracts_incremental_days=int(args.contracts_incremental_days),
        )
    except Exception as exc:  # noqa: BLE001
        print(f"FATAL: {exc}", file=sys.stderr)
        traceback.print_exc()
        return EXIT_TECH

    print(
        json.dumps(
            {
                "cycle_id": report.cycle_id,
                "collection_id": report.collection_id,
                "exit_code": report.exit_code,
                "duration_seconds": report.duration_seconds,
                "products": {
                    k: report.products.get(k)
                    for k in ("executive_md", "excel", "manifest", "pdf_status")
                },
                "intelligence_counts": report.intelligence.get("counts"),
                "runs": [
                    {"source": r.get("source"), "terminal_status": r.get("terminal_status")}
                    for r in report.runs
                ],
                "human_accept": report.human_accept.get("status"),
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        )
    )
    return int(report.exit_code)


if __name__ == "__main__":
    sys.exit(main())
